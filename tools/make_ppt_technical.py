#!/usr/bin/env python3
"""Build a TECHNICAL slide deck for the house cleaner robot.

Same simple, visual presentation style as the intro deck, but every concept is
named with its real ROS 2 / Nav2 term: topics, actions, services, QoS, MPPI,
NavFn, SLAM, costmap layers, lifecycle states, and so on.

Run with the conda python (python-pptx lives there, not in the ROS system
interpreter):

    ~/anaconda3/bin/python tools/make_ppt_technical.py

Output: docs/House_Cleaner_Robot_Technical.pptx
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- palette
NAVY = RGBColor(0x14, 0x2B, 0x50)
BLUE = RGBColor(0x1F, 0x6F, 0xB2)
DARK = RGBColor(0x0E, 0x1E, 0x38)
GREEN = RGBColor(0x17, 0x8F, 0x54)
ORANGE = RGBColor(0xE0, 0x7B, 0x1B)
RED = RGBColor(0xC0, 0x39, 0x2B)
PURPLE = RGBColor(0x6B, 0x4C, 0xA8)
GREY = RGBColor(0x5A, 0x64, 0x72)
LIGHT = RGBColor(0xF2, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xFF, 0xF4, 0xE0)
MONO = "Consolas"

SW = Inches(13.333)
SH = Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]

TOTAL = 16
_n = {"i": 0}


# ---------------------------------------------------------------- helpers
def blank():
    return prs.slides.add_slide(BLANK)


def box(slide, x, y, w, h, fill=None, line=None, lw=1.25,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12):
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except (IndexError, KeyError):
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def text(slide, x, y, w, h, runs, size=18, color=NAVY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0,
         italic=False, font="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    if isinstance(runs, str):
        runs = [(runs, {})]
    first = True
    for item in runs:
        if isinstance(item, str):
            item = (item, {})
        s, opt = item
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = opt.get("align", align)
        para.line_spacing = opt.get("spacing", spacing)
        run = para.add_run()
        run.text = s
        f = run.font
        f.size = Pt(opt.get("size", size))
        f.bold = opt.get("bold", bold)
        f.italic = opt.get("italic", italic)
        f.name = opt.get("font", font)
        f.color.rgb = opt.get("color", color)
    return tb


def slide_base(title, num, kicker=None):
    s = blank()
    box(s, Emu(0), Emu(0), SW, Inches(0.10), fill=BLUE,
        shape=MSO_SHAPE.RECTANGLE)
    text(s, Inches(0.62), Inches(0.28), Inches(10.9), Inches(0.60),
         title, size=27, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        text(s, Inches(0.64), Inches(0.88), Inches(11.6), Inches(0.34),
             kicker, size=13, color=GREY, italic=True)
    text(s, Inches(12.25), Inches(6.94), Inches(0.85), Inches(0.28),
         f"{num}/{TOTAL}", size=11, color=GREY, align=PP_ALIGN.RIGHT)
    return s


def card(slide, x, y, w, h, n, title, body, accent=BLUE, tsize=16,
         bsize=12, mono=False):
    box(slide, x, y, w, h, fill=WHITE, line=RGBColor(0xD3, 0xDC, 0xE8))
    box(slide, x, y, Inches(0.075), h, fill=accent, shape=MSO_SHAPE.RECTANGLE)
    text(slide, x + Inches(0.24), y + Inches(0.14), w - Inches(0.46),
         Inches(0.38), f"{n}.  {title}" if n else title,
         size=tsize, color=accent, bold=True, anchor=MSO_ANCHOR.TOP)
    text(slide, x + Inches(0.24), y + Inches(0.66), w - Inches(0.46),
         h - Inches(0.80), body, size=bsize,
         color=RGBColor(0x33, 0x3D, 0x49), anchor=MSO_ANCHOR.TOP,
         spacing=1.06, font=MONO if mono else "Calibri")


def footer(slide, msg, color=LIGHT, tcolor=NAVY):
    box(slide, Inches(0.62), Inches(6.44), Inches(11.48), Inches(0.58),
        fill=color, line=None)
    text(slide, Inches(0.86), Inches(6.44), Inches(11.0), Inches(0.58),
         msg, size=13, color=tcolor, bold=True)


def arrow(slide, x, y, w, h, color=BLUE):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def chain_slot(slide, x, y, w, h, top, bottom, accent=BLUE):
    """A labelled link in a pipeline: title + mono topic name."""
    box(slide, x, y, w, h, fill=LIGHT, line=accent, lw=1.4)
    text(slide, x + Inches(0.06), y + Inches(0.09), w - Inches(0.12),
         Inches(0.34), top, size=12, color=accent, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if bottom:
        text(slide, x + Inches(0.06), y + Inches(0.44), w - Inches(0.12),
             Inches(0.42), bottom, size=10, color=GREY, font=MONO,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def code(slide, x, y, w, h, lines, fill=DARK, size=12):
    box(slide, x, y, w, h, fill=fill, line=None)
    body = [("%s" % ln, {"size": size, "color": RGBColor(0x8F, 0xE3, 0xB0)
                         if ln.strip().startswith(("#", "$"))
                         else RGBColor(0xE6, 0xED, 0xF6), "font": MONO})
            for ln in lines]
    text(slide, x + Inches(0.20), y + Inches(0.14), w - Inches(0.40),
         h - Inches(0.28), body, anchor=MSO_ANCHOR.TOP, spacing=1.12)


# ================================================================ 1 TITLE
s = blank()
box(s, Emu(0), Emu(0), SW, Inches(2.62), fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
box(s, Emu(0), Inches(2.62), SW, Inches(0.09), fill=BLUE,
    shape=MSO_SHAPE.RECTANGLE)
text(s, Inches(0.9), Inches(0.72), Inches(11.5), Inches(0.9),
     "House Cleaner Robot", size=46, color=WHITE, bold=True)
text(s, Inches(0.9), Inches(1.66), Inches(11.5), Inches(0.58),
     "Autonomous coverage cleaning \u2014 ROS 2 + Nav2 + Gazebo",
     size=21, color=RGBColor(0xBF, 0xD4, 0xEE))
text(s, Inches(0.9), Inches(3.06), Inches(11.5), Inches(0.5),
     "Architecture, control stack, and the failure modes that were measured "
     "and fixed", size=17, color=GREY, italic=True)

# interface strip: the four contracts that define the system
ifaces = [
    ("Action", "NavigateToPose", "/navigate_to_pose", BLUE),
    ("Service", "CreepToDock", "/creep_to_dock", GREEN),
    ("Topic", "BatteryState", "/battery_state", ORANGE),
    ("Topic", "Bool (latched)", "/dock/seated", PURPLE),
]
x = Inches(0.95)
cw = Inches(2.86)
for kind, name, topic, c in ifaces:
    box(s, x, Inches(3.98), cw, Inches(1.62), fill=LIGHT,
        line=RGBColor(0xC5, 0xD2, 0xE2))
    text(s, x + Inches(0.14), Inches(4.10), cw - Inches(0.28), Inches(0.32),
         kind.upper(), size=11, color=c, bold=True, align=PP_ALIGN.CENTER)
    text(s, x + Inches(0.14), Inches(4.48), cw - Inches(0.28), Inches(0.36),
         name, size=15, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text(s, x + Inches(0.14), Inches(4.88), cw - Inches(0.28), Inches(0.40),
         topic, size=11, color=GREY, font=MONO, align=PP_ALIGN.CENTER)
    x += Inches(3.00)

footer(s, "Every subsystem communicates only through these contracts \u2014 "
          "which is why mode:=sim and mode:=real run the same supervisor.")
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "Framing: this is a contract-driven autonomy stack. The mission layer "
    "never imports simulator code, which is what makes the sim-to-real swap "
    "a launch argument rather than a rewrite."
)

# ================================================================ 2 SCOPE
s = slide_base("Scope & design principles", 2,
               "What this system is, and the constraints it was built under")
items = [
    ("Hardware-agnostic mission layer",
     "The supervisor talks only to Nav2 actions, ROS topics and two custom "
     "services. battery_sim and battery_hw are interchangeable backends "
     "behind the same /battery_state + /dock/state contract."),
    ("Portability by detection",
     "No hardcoded distro, username or path. ros-distro.sh resolves "
     "ROS_DISTRO, CONTAINER_NAME and IMAGE_NAME from what is installed; "
     "every script sources it."),
    ("Fail-fast, then skip",
     "A bad coverage cell costs one bounded timeout (45 s) or a detector "
     "trip (~15 s \u2014 NAV_SILENT). It never blocks the mission."),
    ("Single safety gate",
     "Both Nav2 output and the docking creep are published to /cmd_vel_nav, "
     "so velocity_smoother and collision_monitor gate every command."),
]
x = Inches(0.62)
for i, (t, b) in enumerate(items):
    card(s, x, Inches(1.44), Inches(2.92), Inches(4.30), i + 1, t, b,
         accent=[BLUE, PURPLE, ORANGE, GREEN][i], tsize=14, bsize=11.5)
    x += Inches(3.07)
footer(s, "Stack: ROS 2 (rclpy) \u00b7 Gazebo \u00b7 slam_toolbox \u00b7 Nav2 "
          "(MPPI + NavFn) \u00b7 ros_gz_bridge \u00b7 Foxglove \u00b7 Docker")
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "The four principles are the review checklist. Any proposed change that "
    "breaks one of them (e.g. importing simulator internals into the "
    "supervisor) is rejected."
)

# ================================================================ 3 NODEGRAPH
s = slide_base("Node graph & topic contracts", 3,
               "Processes, the interfaces between them, and the direction of data")
# left column: sim + perception
box(s, Inches(0.62), Inches(1.40), Inches(3.55), Inches(1.32), fill=WHITE,
    line=RGBColor(0xD3, 0xDC, 0xE8))
box(s, Inches(0.62), Inches(1.40), Inches(0.075), Inches(1.32), fill=GREY,
    shape=MSO_SHAPE.RECTANGLE)
text(s, Inches(0.88), Inches(1.52), Inches(3.2), Inches(0.32),
     "Simulation", size=15, color=GREY, bold=True)
text(s, Inches(0.88), Inches(1.90), Inches(3.2), Inches(0.72),
     "gz sim (house_room.world)\nros_gz_bridge \u2192 /scan /odom /tf /clock",
     size=11, color=RGBColor(0x33, 0x3D, 0x49), font=MONO, spacing=1.15)

box(s, Inches(0.62), Inches(2.86), Inches(3.55), Inches(1.32), fill=WHITE,
    line=RGBColor(0xD3, 0xDC, 0xE8))
box(s, Inches(0.62), Inches(2.86), Inches(0.075), Inches(1.32), fill=BLUE,
    shape=MSO_SHAPE.RECTANGLE)
text(s, Inches(0.88), Inches(2.98), Inches(3.2), Inches(0.32),
     "Localization", size=15, color=BLUE, bold=True)
text(s, Inches(0.88), Inches(3.36), Inches(3.2), Inches(0.72),
     "slam_toolbox (async)\n\u2192 /map  +  map\u2192odom TF",
     size=11, color=RGBColor(0x33, 0x3D, 0x49), font=MONO, spacing=1.15)

box(s, Inches(0.62), Inches(4.32), Inches(3.55), Inches(1.86), fill=WHITE,
    line=RGBColor(0xD3, 0xDC, 0xE8))
box(s, Inches(0.62), Inches(4.32), Inches(0.075), Inches(1.86), fill=PURPLE,
    shape=MSO_SHAPE.RECTANGLE)
text(s, Inches(0.88), Inches(4.44), Inches(3.2), Inches(0.32),
     "Application nodes", size=15, color=PURPLE, bold=True)
text(s, Inches(0.88), Inches(4.82), Inches(3.2), Inches(1.24),
     "mission_supervisor\nbattery_sim  (or battery_hw)\ndocking_controller",
     size=11, color=RGBColor(0x33, 0x3D, 0x49), font=MONO, spacing=1.20)

# right column: nav2 pipeline
box(s, Inches(4.42), Inches(1.40), Inches(8.30), Inches(4.78), fill=LIGHT,
    line=RGBColor(0xC5, 0xD2, 0xE2))
text(s, Inches(4.66), Inches(1.50), Inches(7.8), Inches(0.32),
     "Nav2 lifecycle stack (brought up by retry loop, not lifecycle_manager)",
     size=12, color=GREY, bold=True)

rows = [
    ("bt_navigator", "NavigateToPose  \u2192  behavior tree", BLUE),
    ("planner_server", "NavFn Dijkstra  \u2192  /plan", BLUE),
    ("controller_server", "MPPI  \u2192  /cmd_vel_nav", GREEN),
    ("behavior_server", "spin / backup / drive_on_heading", GREEN),
    ("velocity_smoother", "/cmd_vel_nav  \u2192  /cmd_vel_smoothed", ORANGE),
    ("collision_monitor", "/cmd_vel_smoothed  \u2192  /cmd_vel", RED),
]
y = Inches(1.92)
for name, role, c in rows:
    box(s, Inches(4.66), y, Inches(7.82), Inches(0.56), fill=WHITE,
        line=RGBColor(0xD3, 0xDC, 0xE8))
    box(s, Inches(4.66), y, Inches(0.07), Inches(0.56), fill=c,
        shape=MSO_SHAPE.RECTANGLE)
    text(s, Inches(4.88), y, Inches(2.45), Inches(0.56), name, size=11.5,
         color=c, bold=True, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(7.45), y, Inches(4.9), Inches(0.56), role, size=11,
         color=RGBColor(0x33, 0x3D, 0x49), font=MONO,
         anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.64)

text(s, Inches(4.66), Inches(5.82), Inches(7.82), Inches(0.30),
     "/cmd_vel  \u2192  ros_gz_bridge  \u2192  gz DiffDrive plugin",
     size=11, color=GREY, font=MONO, align=PP_ALIGN.CENTER)
footer(s, "Config is bind-mounted, so src/ and YAML edits need no rebuild \u2014 "
          "relaunch picks them up (symlink-install).")
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "Walk the chain left to right. The key point: controller_server is "
    "remapped to cmd_vel_nav so its raw output never reaches the robot "
    "directly."
)

# ================================================================ 4 VELCHAIN
s = slide_base("The velocity chain", 4,
               "Four stages, one message type \u2014 and the bug that breaks it")
chain = [
    ("controller / behavior\n/ docking", "/cmd_vel_nav", GREEN),
    ("velocity_smoother", "/cmd_vel_smoothed", ORANGE),
    ("collision_monitor", "/cmd_vel", RED),
    ("ros_gz_bridge", "gz.msgs.Twist", GREY),
]
x = Inches(0.80)
cw = Inches(2.80)
for i, (top, bottom, c) in enumerate(chain):
    chain_slot(s, x, Inches(1.50), cw, Inches(1.00),
               top.replace("\n", " "), bottom, accent=c)
    if i < 3:
        arrow(s, x + cw + Inches(0.03), Inches(1.86), Inches(0.22),
              Inches(0.26), color=RGBColor(0x9F, 0xB4, 0xCC))
    x += cw + Inches(0.30)

text(s, Inches(0.62), Inches(2.78), Inches(12.1), Inches(0.34),
     "Rates: 20 Hz controller \u00b7 20 Hz smoother \u00b7 20 Hz monitor \u00b7 "
     "collision_monitor source_timeout 3.0 s", size=12.5, color=GREY)

box(s, Inches(0.62), Inches(3.26), Inches(5.95), Inches(2.44), fill=WHITE,
    line=RED, lw=1.5)
text(s, Inches(0.88), Inches(3.40), Inches(5.45), Inches(0.34),
     "Failure mode: type mismatch", size=15, color=RED, bold=True)
text(s, Inches(0.88), Inches(3.80), Inches(5.45), Inches(1.80),
     "Newer distros default enable_stamped_cmd_vel: true (TwistStamped). "
     "velocity_smoother, collision_monitor and the gz bridge all speak plain "
     "geometry_msgs/Twist, so the smoother receives nothing and the robot "
     "never moves \u2014 with no error logged anywhere.", size=11.5,
     color=RGBColor(0x33, 0x3D, 0x49))

box(s, Inches(6.78), Inches(3.26), Inches(5.94), Inches(2.44), fill=WHITE,
    line=GREEN, lw=1.5)
text(s, Inches(7.04), Inches(3.40), Inches(5.44), Inches(0.34),
     "Fix", size=15, color=GREEN, bold=True)
code(s, Inches(7.04), Inches(3.78), Inches(5.44), Inches(1.72), [
    "controller_server:",
    "  ros__parameters:",
    "    enable_stamped_cmd_vel: false",
    "behavior_server:   # same",
    "velocity_smoother: # same",
    "collision_monitor: # same",
], size=11)
footer(s, "Rule: every stage in the chain must agree on the message type. "
          "Verified with ros2 topic hz on each stage.")
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "This is the highest-value slide for anyone porting the stack to a new "
    "distro. Symptom is 'no motion, no errors'."
)

# ================================================================ 5 BOOTGATE
s = slide_base("Startup & the lifecycle gate", 5,
               "Why goals are not sent until three nodes report ACTIVE")
steps = [
    ("1. gz sim + bridge", "Gazebo starts, bridge republishes "
     "/clock /scan /odom /tf", GREY),
    ("2. slam_toolbox", "configure + activate via retry loop (120\u00d71 s). "
     "No Bond server, so no lifecycle_manager.", BLUE),
    ("3. Nav2 nodes", "8 nodes, each configure then activate, retried. "
     "Ends with the READY marker.", BLUE),
    ("4. Supervisor wait", "Concurrent wait on /map AND the "
     "NavigateToPose action server, 180 s cap.", ORANGE),
    ("5. ACTIVE gate", "Polls get_state on bt_navigator, velocity_smoother, "
     "collision_monitor; requires id == 3.", GREEN),
    ("6. Plan + first goal", "Coverage goals built from the live map; "
     "goal 1 gets a 90 s timeout.", GREEN),
]
y = Inches(1.42)
for i, (t, b, c) in enumerate(steps):
    box(s, Inches(0.62), y, Inches(12.1), Inches(0.74), fill=WHITE,
        line=RGBColor(0xD3, 0xDC, 0xE8))
    box(s, Inches(0.62), y, Inches(0.075), Inches(0.74), fill=c,
        shape=MSO_SHAPE.RECTANGLE)
    text(s, Inches(0.90), y, Inches(3.05), Inches(0.74), t, size=13,
         color=c, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(4.10), y, Inches(8.42), Inches(0.74), b, size=11.5,
         color=RGBColor(0x33, 0x3D, 0x49), anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.82)
footer(s, "Cause of the gate: the NavigateToPose server is advertised before "
          "the node is ACTIVE, and velocity_smoother/collision_monitor "
          "activate LAST.")
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "Without step 5 the controller accepts a goal while the velocity chain "
    "is dead. That looks like a stall, which fed a watchdog restart loop."
)

# ================================================================ 6 COVERAGE
s = slide_base("Coverage planning: boustrophedon", 6,
               "Pure-Python geometry in house_cleaner_core, unit-testable "
               "without ROS")
box(s, Inches(0.62), Inches(1.44), Inches(6.5), Inches(4.55), fill=LIGHT,
    line=RGBColor(0xC5, 0xD2, 0xE2))
text(s, Inches(0.62), Inches(1.50), Inches(6.5), Inches(0.30),
     "row endpoints become navigation goals", size=11, color=GREY,
     italic=True, align=PP_ALIGN.CENTER)
px, py, pw, ph = Inches(1.00), Inches(1.95), Inches(5.75), Inches(3.70)
box(s, px, py, pw, ph, fill=WHITE, line=NAVY, lw=2,
    shape=MSO_SHAPE.RECTANGLE)
rows_, step = 7, ph / 7.0
for r in range(rows_):
    yy = py + Emu(int(step * (r + 0.5)))
    if r % 2 == 0:
        x1, x2 = px + Inches(0.15), px + pw - Inches(0.15)
    else:
        x1, x2 = px + pw - Inches(0.15), px + Inches(0.15)
    a = min(x1, x2)
    sh = s.shapes.add_shape(MSO_SHAPE.LEFT_RIGHT_ARROW, a,
                            yy - Inches(0.035),
                            abs(x2 - x1), Inches(0.07))
    sh.fill.solid()
    sh.fill.fore_color.rgb = BLUE
    sh.line.fill.background()
    sh.shadow.inherit = False
    cx = (x2 - Inches(0.30)) if r % 2 == 0 else (x1 + Inches(0.10))
    cy = yy + step * 0.26 if r % 2 == 0 else yy - step * 0.40
    box(s, cx, cy, Inches(0.18), Inches(0.18), fill=GREEN,
        shape=MSO_SHAPE.OVAL)
text(s, Inches(0.62), Inches(5.72), Inches(6.5), Inches(0.28),
     "two goals per row, alternating yaw (0 / \u03c0)",
     size=11, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

acts = [
    ("coverage_bounds(grid, 0.55)",
     "Coverage rectangle from the SLAM map window, inset 0.55 m to clear the "
     "0.30 m inflation band."),
    ("build_boustrophedon(bounds, 0.45)",
     "Rows 0.45 m apart; each row contributes its two endpoints, with "
     "alternating yaw."),
    ("snap_to_free_row(grid, x, y)",
     "Plan-time snap: nearest cell with 0 \u2264 occupancy < 50; unknown (-1) "
     "is never a goal."),
    ("revalidate_goal(...) before every goal",
     "Re-checks the newest costmap, sweeps \u00b12 m along the row, or drops "
     "the goal with a reason."),
]
y = Inches(1.44)
for i, (t, b) in enumerate(acts):
    card(s, Inches(7.32), y, Inches(5.40), Inches(1.10), 0, t, b,
         accent=GREEN, tsize=12.5, bsize=11, mono=False)
    y += Inches(1.18)
footer(s, "Static furniture AABBs are excluded twice: at plan time "
          "(in_furniture_keepout) and again at goal time.")
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "geometry.py is deliberately ROS-free so the coverage maths can be "
    "unit-tested headlessly. Function names are the real API."
)

# ================================================================ 7 PLANNER
s = slide_base("Planning & control: NavFn + MPPI", 7,
               "Global grid search, then a sampling-based local controller")
box(s, Inches(0.62), Inches(1.44), Inches(6.05), Inches(2.30), fill=WHITE,
    line=BLUE, lw=1.5)
text(s, Inches(0.88), Inches(1.56), Inches(5.55), Inches(0.32),
     "planner_server \u00b7 NavFn", size=16, color=BLUE, bold=True)
text(s, Inches(0.88), Inches(1.98), Inches(5.55), Inches(1.62),
     "Dijkstra expansion over the global costmap, then backtrack. "
     "allow_unknown: true (early SLAM map is mostly unknown). "
     "tolerance: 0.30 \u2014 must stay \u2264 the goal checker's xy tolerance "
     "or the goal can never be satisfied.", size=11.5,
     color=RGBColor(0x33, 0x3D, 0x49))

box(s, Inches(6.88), Inches(1.44), Inches(5.84), Inches(2.30), fill=WHITE,
    line=GREEN, lw=1.5)
text(s, Inches(7.14), Inches(1.56), Inches(5.34), Inches(0.32),
     "controller_server \u00b7 MPPI", size=16, color=GREEN, bold=True)
text(s, Inches(7.14), Inches(1.98), Inches(5.34), Inches(1.62),
     "Samples noisy velocity rollouts, scores them with weighted critics, and "
     "averages the best. batch_size 600 \u00b7 time_steps 56 \u00d7 model_dt "
     "0.05 s (2.8 s horizon) \u00b7 temperature 0.5 \u00b7 gamma 0.015 \u00b7 "
     "20 Hz.", size=11.5, color=RGBColor(0x33, 0x3D, 0x49))

text(s, Inches(0.62), Inches(3.86), Inches(12.1), Inches(0.30),
     "MPPI critics and weights", size=14, color=NAVY, bold=True)
crit = [("PathAlign 14.0", BLUE), ("PathFollow 7.0", BLUE),
        ("Goal 5.0", GREEN), ("Constraint 4.0", GREEN),
        ("Cost 3.5", ORANGE), ("GoalAngle 3.0", ORANGE),
        ("PathAngle 2.0", PURPLE), ("PreferForward 0.0", RED)]
x = Inches(0.62)
cw = Inches(1.46)
for name, c in crit:
    box(s, x, Inches(4.24), cw, Inches(0.60), fill=LIGHT, line=c, lw=1.2)
    text(s, x + Inches(0.04), Inches(4.24), cw - Inches(0.08), Inches(0.60),
         name, size=10.5, color=c, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    x += Inches(1.54)

box(s, Inches(0.62), Inches(5.02), Inches(12.1), Inches(1.18), fill=CREAM,
    line=RGBColor(0xF0, 0xD2, 0x9A))
text(s, Inches(0.90), Inches(5.02), Inches(11.54), Inches(1.18),
     "Consequence of PreferForward = 0.0: reverse is free, which is what lets "
     "the robot back out of furniture corners. The trade-off is that the cost "
     "gradient can vanish near a goal, so MPPI may stop ~0.28 m short \u2014 "
     "hence xy_goal_tolerance 0.30 and the NAV_SILENT detector.",
     size=11.5, color=RGBColor(0x8A, 0x54, 0x0A), anchor=MSO_ANCHOR.MIDDLE)
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "Critic weights are the actual YAML values. The 0.0 on PreferForward is "
    "a deliberate choice with a documented trade-off."
)

# ================================================================ 8 SLAM
s = slide_base("SLAM & costmaps", 8,
               "slam_toolbox scan matching, and the costmap that feeds NavFn")
box(s, Inches(0.62), Inches(1.44), Inches(6.05), Inches(4.55), fill=WHITE,
    line=BLUE, lw=1.5)
text(s, Inches(0.88), Inches(1.56), Inches(5.55), Inches(0.32),
     "slam_toolbox (async)", size=16, color=BLUE, bold=True)
text(s, Inches(0.88), Inches(1.98), Inches(5.55), Inches(3.90),
     [("resolution 0.05 m, mode: mapping, scan_period 0.1 s\n\n",
       {"size": 11.5}),
      ("minimum_travel_distance 0.15 m\nminimum_travel_heading 0.20 rad\n\n",
       {"size": 11.5, "font": MONO}),
      ("Both must stay BELOW the 0.48 m undock move. At the original 0.5 m "
       "the final undock never triggered a scan match, so map\u2192odom stayed "
       "stale, the robot's map pose sat inside the dock's lethal cells, and "
       "the planner failed with \u201cFailed to create plan\u201d on every "
       "goal.", {"size": 11.5})],
     anchor=MSO_ANCHOR.TOP, spacing=1.05)

box(s, Inches(6.88), Inches(1.44), Inches(5.84), Inches(4.55), fill=WHITE,
    line=GREEN, lw=1.5)
text(s, Inches(7.14), Inches(1.56), Inches(5.34), Inches(0.32),
     "Costmap layers", size=16, color=GREEN, bold=True)
rows = [
    ("static_layer", "from /map, transient_local subscription"),
    ("obstacle_layer", "/scan, marking + clearing, raycast 3.0 m"),
    ("inflation_layer", "radius 0.30 m, cost_scaling_factor 2.5"),
    ("robot_radius", "0.105 m (Burger)"),
]
y = Inches(2.04)
for name, desc in rows:
    box(s, Inches(7.14), y, Inches(5.34), Inches(0.52), fill=LIGHT)
    text(s, Inches(7.26), y, Inches(2.0), Inches(0.52), name, size=11,
         color=GREEN, bold=True, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(9.30), y, Inches(3.10), Inches(0.52), desc, size=10.5,
         color=RGBColor(0x33, 0x3D, 0x49), anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.60)
text(s, Inches(7.14), Inches(4.70), Inches(5.34), Inches(1.20),
     "Critical difference: global_costmap uses track_unknown_space: true, "
     "but local_costmap uses false. With true, the rolling window stays ~97% "
     "unknown, MPPI treats it as blocked and the robot never moves.",
     size=11, color=RGBColor(0x33, 0x3D, 0x49), anchor=MSO_ANCHOR.TOP)
footer(s, "inflation_radius 0.30 m is why coverage goals need a 0.55 m "
          "margin \u2014 goals must not land inside the inflation band.")
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "Two parameters caused multi-day bugs: the SLAM travel gate and "
    "track_unknown_space on the local costmap."
)

# ================================================================ 9 MISSION
s = slide_base("Mission state machine", 9,
               "mission_supervisor: the top-level FSM and its guards")
states = [("INIT", GREY), ("CLEANING", BLUE), ("RETURNING", ORANGE),
          ("DOCKING", PURPLE), ("CHARGING", GREEN), ("UNDOCKING", ORANGE),
          ("DONE", GREY)]
x = Inches(0.62)
cw = Inches(1.62)
for i, (st, c) in enumerate(states):
    box(s, x, Inches(1.46), cw, Inches(0.72), fill=WHITE, line=c, lw=1.5)
    text(s, x + Inches(0.04), Inches(1.46), cw - Inches(0.08), Inches(0.72),
         st, size=12, color=c, bold=True, font=MONO,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < len(states) - 1:
        arrow(s, x + cw + Inches(0.02), Inches(1.70), Inches(0.16),
              Inches(0.24), color=RGBColor(0xA9, 0xBA, 0xCE))
    x += cw + Inches(0.20)

text(s, Inches(0.62), Inches(2.34), Inches(12.1), Inches(0.30),
     "Transitions", size=14, color=NAVY, bold=True)
trans = [
    ("CLEANING \u2192 RETURNING", "battery \u2264 low_threshold (40%) \u2014 "
     "watchdog timer cancels the in-flight goal"),
    ("RETURNING \u2192 DOCKING", "an approach pose reached \u2192 "
     "/creep_to_dock service"),
    ("DOCKING \u2192 CHARGING", "response.seated == true \u2192 "
     "/dock/command CMD_START"),
    ("CHARGING \u2192 UNDOCKING", "percentage \u2265 95% \u2192 CMD_STOP "
     "\u2192 /undock Trigger"),
    ("UNDOCKING \u2192 CLEANING", "verified reverse motion \u2265 0.24 m"),
    ("any \u2192 DONE", "battery \u2264 1% and not seated (hard abort)"),
]
y = Inches(2.70)
for t, b in trans:
    box(s, Inches(0.62), y, Inches(12.1), Inches(0.52), fill=LIGHT)
    text(s, Inches(0.86), y, Inches(3.5), Inches(0.52), t, size=11.5,
         color=BLUE, bold=True, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(4.55), y, Inches(7.95), Inches(0.52), b, size=11,
         color=RGBColor(0x33, 0x3D, 0x49), anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.60)
footer(s, "mission.loop := true restarts coverage after the final dock \u2014 "
          "used for soak testing and demos.")
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "The FSM is the spine of the supervisor. Each transition has a single "
    "observable condition, which is what makes it testable."
)

# ================================================================ 10 TIME
s = slide_base("Goal timeouts & the result policy", 10,
               "Every reachable outcome of a NavigateToPose call has a handler")
to = [
    ("CLEANING goal 1", "90 s", "cold start: map/TF still maturing"),
    ("CLEANING goals 2..n", "45 s", "normal goal finishes in 4\u201313 s"),
    ("RETURNING", "90 s", "capped by the remaining return budget"),
    ("mission.return_budget", "180 s", "wall clock for the whole return phase"),
    ("progress_checker", "0.05 m / 90 s", "abort when no progress"),
    ("charge wait", "300 s", "bounded; exits early if dock lost"),
]
x = Inches(0.62)
cw = Inches(1.98)
for i, (name, val, why) in enumerate(to):
    cx = x + Inches(2.04) * i
    box(s, cx, Inches(1.42), cw, Inches(1.68), fill=WHITE,
        line=RGBColor(0xD3, 0xDC, 0xE8))
    box(s, cx, Inches(1.42), cw, Inches(0.08), fill=BLUE,
        shape=MSO_SHAPE.RECTANGLE)
    text(s, cx + Inches(0.12), Inches(1.58), cw - Inches(0.24), Inches(0.34),
         name, size=11.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text(s, cx + Inches(0.12), Inches(1.94), cw - Inches(0.24), Inches(0.34),
         val, size=15, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    text(s, cx + Inches(0.12), Inches(2.32), cw - Inches(0.24), Inches(0.70),
         why, size=10, color=GREY, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.TOP)

text(s, Inches(0.62), Inches(3.26), Inches(12.1), Inches(0.30),
     "Result handling", size=14, color=NAVY, bold=True)
res = [
    ("SUCCEEDED", "advance", GREEN),
    ("ABORTED err=102", "TF lag \u2014 settle 3 s, retry once", ORANGE),
    ("ABORTED other", "planner proved it unreachable \u2014 skip", RED),
    ("TIMEOUT", "revalidate, retry once, then skip", ORANGE),
    ("NAV_SILENT", "MPPI frozen \u2014 skip immediately", RED),
    ("VEL_STALL", "settle, one retry, then skip", ORANGE),
    ("RETURN_BUDGET", "stop return attempts", RED),
    ("BATTERY_EMPTY", "abort mission cycle", RED),
    ("CANCELED", "low-battery cancel \u2192 recharge cycle", BLUE),
]
y = Inches(3.62)
for i, (st, act, c) in enumerate(res):
    cx = Inches(0.62) + Inches(4.08) * (i % 3)
    cy = y + Inches(0.62) * (i // 3)
    box(s, cx, cy, Inches(3.92), Inches(0.54), fill=WHITE,
        line=RGBColor(0xD3, 0xDC, 0xE8))
    box(s, cx, cy, Inches(0.07), Inches(0.54), fill=c,
        shape=MSO_SHAPE.RECTANGLE)
    text(s, cx + Inches(0.18), cy, Inches(1.45), Inches(0.54), st, size=10.5,
         color=c, bold=True, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    text(s, cx + Inches(1.66), cy, Inches(2.16), Inches(0.54), act, size=10,
         color=RGBColor(0x33, 0x3D, 0x49), anchor=MSO_ANCHOR.MIDDLE)
footer(s, "No status reaches the loop without an explicit handler \u2014 that "
          "is what keeps a bad cell bounded instead of fatal.")
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "Statuses are the real string values returned by send_goal(). "
    "NAV_SILENT and VEL_STALL are project-specific detectors layered on top "
    "of Nav2's statuses."
)

# ================================================================ 11 DETECT
s = slide_base("Fault detectors & watchdog layers", 11,
               "Three tiers: node timer, mission detectors, host supervisor")
box(s, Inches(0.62), Inches(1.42), Inches(4.06), Inches(4.55), fill=WHITE,
    line=BLUE, lw=1.5)
box(s, Inches(0.62), Inches(1.42), Inches(4.06), Inches(0.08), fill=BLUE,
    shape=MSO_SHAPE.RECTANGLE)
text(s, Inches(0.84), Inches(1.58), Inches(3.6), Inches(0.32),
     "Tier 1 \u00b7 in-node timer", size=14, color=BLUE, bold=True)
text(s, Inches(0.84), Inches(2.00), Inches(3.62), Inches(3.80),
     [("create_timer(0.5, _watchdog_cb):\n", {"size": 11, "font": MONO}),
      ("\u2022 battery \u2264 40% \u2192 cancel_goal_async()\n"
       "\u2022 battery \u2264 1% \u2192 hard cancel\n"
       "\u2022 VEL_STALL \u2192 cancel\n"
       "\u2022 return budget exceeded \u2192 cancel\n\n",
       {"size": 11}),
      ("Runs while spin_until_future_complete() renders the node, so it "
       "fires even mid-action.", {"size": 11})],
     anchor=MSO_ANCHOR.TOP, spacing=1.08)

box(s, Inches(4.84), Inches(1.42), Inches(4.06), Inches(4.55), fill=WHITE,
    line=ORANGE, lw=1.5)
box(s, Inches(4.84), Inches(1.42), Inches(4.06), Inches(0.08), fill=ORANGE,
    shape=MSO_SHAPE.RECTANGLE)
text(s, Inches(5.06), Inches(1.58), Inches(3.6), Inches(0.32),
     "Tier 2 \u00b7 mission detectors", size=14, color=ORANGE, bold=True)
text(s, Inches(5.06), Inches(2.00), Inches(3.62), Inches(3.80),
     [("_nav_abort_reason(), polled at 100 ms:\n\n", {"size": 11}),
      ("VEL_STALL\n", {"size": 11, "font": MONO, "bold": True}),
      ("/cmd_vel_nav fresh and v \u2265 0.02, but /cmd_vel silent > "
       "stall_hold (8 s) and robot still\n\n", {"size": 10.5}),
      ("NAV_SILENT\n", {"size": 11, "font": MONO, "bold": True}),
      ("controller alive but v,w < 0.02 for \u2265 15 s\n\n",
       {"size": 10.5}),
      ("Reason strings are returned as the goal status itself.",
       {"size": 10.5})],
     anchor=MSO_ANCHOR.TOP, spacing=1.06)

box(s, Inches(9.06), Inches(1.42), Inches(3.66), Inches(4.55), fill=WHITE,
    line=RED, lw=1.5)
box(s, Inches(9.06), Inches(1.42), Inches(3.66), Inches(0.08), fill=RED,
    shape=MSO_SHAPE.RECTANGLE)
text(s, Inches(9.28), Inches(1.58), Inches(3.2), Inches(0.32),
     "Tier 3 \u00b7 host watchdog", size=14, color=RED, bold=True)
text(s, Inches(9.28), Inches(2.00), Inches(3.22), Inches(3.80),
     [("watchdog.sh, every 15 s:\n\n", {"size": 11}),
      ("\u2022 container down\n\u2022 no mission log 180 s\n"
       "\u2022 \u22654 stall lines / 100 s with no progress\n"
       "\u2022 battery empty, not seated\n"
       "\u2022 /cmd_vel_nav \u2265 2 Hz but /cmd_vel < 1 Hz (twice)\n\n",
       {"size": 10.5}),
      ("Soft: dump + cancel_nav + bounce smoother/monitor. "
       "2 in a row \u2192 full stack restart.", {"size": 10.5})],
     anchor=MSO_ANCHOR.TOP, spacing=1.06)
footer(s, "Escalation is deliberate: local cancel \u2192 bounce the two tail "
          "nodes \u2192 restart the container.")
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "Emphasise the debouncing: the chain probe needs two consecutive hits, "
    "and any progress marker resets the stuck-pattern counter. Unhysteretic "
    "watchdogs caused restart loops."
)

# ================================================================ 12 DOCK
s = slide_base("Docking: laser-guided creep", 12,
               "Laser sector in the base frame, three independent seat latches")
box(s, Inches(0.62), Inches(1.42), Inches(6.05), Inches(2.28), fill=WHITE,
    line=GREEN, lw=1.5)
text(s, Inches(0.88), Inches(1.54), Inches(5.55), Inches(0.30),
     "Front clearance", size=14, color=GREEN, bold=True)
text(s, Inches(0.88), Inches(1.92), Inches(5.55), Inches(0.34),
     "min range within \u00b1creep.sector_deg (12\u00b0) of the forward ray; "
     "index from angle_min, not index 0", size=11,
     color=RGBColor(0x33, 0x3D, 0x49))
seat = [
    ("fwd < 0.13 m", "creep.stop_range \u2014 geometric seat", GREEN),
    ("stall 2.5 s @ fwd < 0.35 m", "odom speed < 0.01 while pressed", ORANGE),
    ("6 samples / 3 s, spread < 0.05 m", "flat-clearance window latch",
     PURPLE),
]
y = Inches(2.40)
for t, b, c in seat:
    box(s, Inches(0.88), y, Inches(5.55), Inches(0.40), fill=LIGHT)
    box(s, Inches(0.88), y, Inches(0.06), Inches(0.40), fill=c,
        shape=MSO_SHAPE.RECTANGLE)
    text(s, Inches(1.04), y, Inches(2.9), Inches(0.40), t, size=10,
         color=c, bold=True, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(3.98), y, Inches(2.36), Inches(0.40), b, size=9.5,
         color=GREY, anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.46)

box(s, Inches(6.88), Inches(1.42), Inches(5.84), Inches(2.28), fill=WHITE,
    line=BLUE, lw=1.5)
text(s, Inches(7.14), Inches(1.54), Inches(5.34), Inches(0.30),
     "Parameters (namespace creep. / undock.)", size=14, color=BLUE, bold=True)
prm = [
    ("creep.speed", "0.04 m/s"),
    ("creep.timeout", "30.0 s"),
    ("creep.max_attempts", "2"),
    ("undock.speed", "\u22120.06 m/s"),
    ("undock.meters", "0.48 m"),
    ("velocity_topic", "/cmd_vel_nav"),
]
y = Inches(1.94)
for name, val in prm:
    box(s, Inches(7.14), y, Inches(5.34), Inches(0.27), fill=LIGHT)
    text(s, Inches(7.26), y, Inches(3.1), Inches(0.27), name, size=9.5,
         color=NAVY, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(10.40), y, Inches(2.0), Inches(0.27), val, size=9.5,
         color=GREY, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.30)

box(s, Inches(0.62), Inches(3.86), Inches(12.1), Inches(2.10), fill=LIGHT,
    line=RGBColor(0xC5, 0xD2, 0xE2))
text(s, Inches(0.88), Inches(3.98), Inches(5.5), Inches(0.30),
     "Interface", size=13, color=NAVY, bold=True)
code(s, Inches(0.88), Inches(4.34), Inches(5.72), Inches(1.48), [
    "house_cleaner_msgs/CreepToDock",
    "  ---",
    "  bool    seated",
    "  float32 front_clearance_m",
    "",
    "std_srvs/Trigger  ->  /undock",
], size=10.5)
text(s, Inches(6.80), Inches(3.98), Inches(5.7), Inches(1.86),
     [("Reporting: /dock/seated (std_msgs/Bool) is published TRANSIENT_LOCAL "
       "so a late or restarted subscriber still learns the state.\n\n",
       {"size": 10.5}),
      ("Verification: undock integrates |v|\u00b7dt and requires "
       "\u2265 0.24 m. If it did not move, seated stays true \u2014 the "
       "backend can never claim to be charging in free space.",
       {"size": 10.5})],
     anchor=MSO_ANCHOR.TOP, spacing=1.06)
footer(s, "The creep is published to /cmd_vel_nav on purpose: the collision "
          "monitor stays the last-resort bumper during docking.")
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "Three independent latches means a single sensor lie cannot fake a seat. "
    "The service interface is the real .srv content."
)

# ================================================================ 13 SAFETY
s = slide_base("Safety chain & collision monitor", 13,
               "Why the stop polygon was removed and what replaced it")
box(s, Inches(0.62), Inches(1.42), Inches(6.05), Inches(2.40), fill=WHITE,
    line=GREEN, lw=1.5)
text(s, Inches(0.88), Inches(1.54), Inches(5.55), Inches(0.30),
     "Kept: PolygonSlow", size=14, color=GREEN, bold=True)
code(s, Inches(0.88), Inches(1.92), Inches(5.55), Inches(1.20), [
    "PolygonSlow:",
    "  type: polygon",
    "  action_type: limit",
    "  points: '[0.35,0.15],[0.35,-0.15],",
    "           [0.12,-0.15],[0.12,0.15]'",
    "  speed_limit_pct: 45.0",
], size=10.5)
text(s, Inches(0.88), Inches(3.20), Inches(5.55), Inches(0.50),
     "Limits speed to 45% inside a 0.12\u20130.35 m front quad. "
     "source_timeout: 3.0 s.", size=10.5,
     color=RGBColor(0x33, 0x3D, 0x49))

box(s, Inches(6.88), Inches(1.42), Inches(5.84), Inches(2.40), fill=WHITE,
    line=RED, lw=1.5)
text(s, Inches(7.14), Inches(1.54), Inches(5.34), Inches(0.30),
     "Removed: PolygonStop", size=14, color=RED, bold=True)
text(s, Inches(7.14), Inches(1.92), Inches(5.34), Inches(1.78),
     "action_type: stop deadlocked the chain. While any scan point sits in "
     "the polygon the monitor forces output to zero, and after "
     "stop_pub_timeout it stops publishing entirely. Because the robot is "
     "stopped, the obstacle never leaves the polygon \u2014 so /cmd_vel stays "
     "silent forever and every later goal looks like VEL_STALL until the "
     "stack is restarted.", size=10.5,
     color=RGBColor(0x33, 0x3D, 0x49))

box(s, Inches(0.62), Inches(3.98), Inches(12.1), Inches(1.98), fill=LIGHT,
    line=RGBColor(0xC5, 0xD2, 0xE2))
text(s, Inches(0.88), Inches(4.10), Inches(11.6), Inches(0.30),
     "Smoother configuration (velocity_smoother)", size=13, color=NAVY,
     bold=True)
pars = [("feedback", "OPEN_LOOP"), ("max_velocity", "[0.65, 0.0, 2.5]"),
        ("max_accel", "[1.5, 0.0, 3.0]"), ("max_decel", "[\u22122.0, 0.0, \u22123.5]"),
        ("max_jerk", "[3.0, 0.0, 5.0]"), ("cmd_vel_timeout", "0.5 s")]
y = Inches(4.48)
for i, (n, v) in enumerate(pars):
    cx = Inches(0.88) + Inches(3.94) * (i % 3)
    cy = y + Inches(0.62) * (i // 3)
    box(s, cx, cy, Inches(3.78), Inches(0.36), fill=WHITE)
    text(s, cx + Inches(0.14), cy, Inches(1.8), Inches(0.36), n, size=10,
         color=NAVY, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    text(s, cx + Inches(1.94), cy, Inches(1.72), Inches(0.36), v, size=10,
         color=GREY, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
footer(s, "Lesson: a fail-safe that stops publishing is not fail-safe for the "
          "rest of the pipeline.")
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "This is a genuinely counter-intuitive result: the safety feature caused "
    "the deadlock. Worth dwelling on for a technical audience."
)

# ================================================================ 14 BATTERY
s = slide_base("Battery model & energy policy", 14,
               "A rate model with hot-tunable parameters")
box(s, Inches(0.62), Inches(1.42), Inches(6.05), Inches(2.42), fill=WHITE,
    line=ORANGE, lw=1.5)
text(s, Inches(0.88), Inches(1.54), Inches(5.55), Inches(0.30),
     "Model (battery_sim, 0.2 s tick)", size=14, color=ORANGE, bold=True)
code(s, Inches(0.88), Inches(1.92), Inches(5.55), Inches(1.80), [
    "if seated and charging and pct < 100:",
    "    pct += charge_rate * dt",
    "elif not docked and speed > 0.03:",
    "    pct -= drain_rate * dt",
    "",
    "# speed from /odom (ground truth),",
    "# /cmd_vel + /cmd_vel_smoothed as backups",
], size=10.5)

box(s, Inches(6.88), Inches(1.42), Inches(5.84), Inches(2.42), fill=WHITE,
    line=BLUE, lw=1.5)
text(s, Inches(7.14), Inches(1.54), Inches(5.34), Inches(0.30),
     "Publishing", size=14, color=BLUE, bold=True)
text(s, Inches(7.14), Inches(1.92), Inches(5.34), Inches(1.80),
     [("/battery_state  sensor_msgs/BatteryState\n", {"size": 11, "font": MONO}),
      ("percentage is a 0\u20131 fraction (ROS standard). Publishing 0\u2013100 "
       "made a real 0.96% read as 96%, which exited the charge loop "
       "immediately.\n\n", {"size": 10.5}),
      ("voltage interpolated 10.0 \u2192 12.6 V linearly with percentage.",
       {"size": 10.5})],
     anchor=MSO_ANCHOR.TOP, spacing=1.06)

text(s, Inches(0.62), Inches(3.98), Inches(12.1), Inches(0.30),
     "Thresholds", size=14, color=NAVY, bold=True)
th = [("low_threshold", "40 %", "enter RETURNING", ORANGE),
      ("charge_target", "95 %", "resume CLEANING", GREEN),
      ("critical", "1 %", "hard mission abort", RED),
      ("drain_rate", "0.20 %/s", "hot-tunable", BLUE),
      ("charge_rate", "0.80 %/s", "hot-tunable", BLUE),
      ("speed_threshold", "0.03 m/s", "motion gate", GREY)]
x = Inches(0.62)
for name, val, why, c in th:
    cx = x
    box(s, cx, Inches(4.36), Inches(1.98), Inches(1.62), fill=WHITE,
        line=RGBColor(0xD3, 0xDC, 0xE8))
    box(s, cx, Inches(4.36), Inches(1.98), Inches(0.07), fill=c,
        shape=MSO_SHAPE.RECTANGLE)
    text(s, cx + Inches(0.10), Inches(4.50), Inches(1.78), Inches(0.50),
         name, size=10, color=c, bold=True, font=MONO,
         align=PP_ALIGN.CENTER)
    text(s, cx + Inches(0.10), Inches(4.98), Inches(1.78), Inches(0.30),
         val, size=13, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text(s, cx + Inches(0.10), Inches(5.32), Inches(1.78), Inches(0.52),
         why, size=9.5, color=GREY, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.TOP)
    x += Inches(2.04)
footer(s, "$ ros2 param set /battery drain_rate 4.0   # accelerate a demo; "
          "0.20 to restore")
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "Both rates are re-read every tick, so a live mission can be sped up "
    "without relaunch. The 0\u20131 percentage convention is a classic "
    "integration bug."
)

# ================================================================ 15 OPS
s = slide_base("Operations: one command, eight checks", 15,
               "Reproducible bring-up, deterministic teardown")
code(s, Inches(0.62), Inches(1.40), Inches(5.95), Inches(0.72), [
    "$ ./run_all.sh          # kill stale \u2192 launch \u2192 verify",
    "$ ./run_all.sh --stop   # kill only",
    "$ ./run_all.sh --verify   # checks only",
], size=11)

checks = [
    "container house_cleaner_<distro> is running",
    "gz-sim-main present (docker top)",
    "rviz2 present (docker top)",
    "log contains 'Nav2 lifecycle bringup complete'",
    "Foxglove WebSocket upgrade returns HTTP 101",
    "/battery_state publishes (ros2 topic echo --once)",
    "watchdog.sh running (pgrep)",
    "monitor_loop.sh running (pgrep)",
]
y = Inches(2.26)
for i, c in enumerate(checks):
    box(s, Inches(0.62), y, Inches(5.95), Inches(0.42), fill=LIGHT)
    box(s, Inches(0.62), y, Inches(0.32), Inches(0.42), fill=GREEN)
    text(s, Inches(0.62), y, Inches(0.32), Inches(0.42), str(i + 1),
         size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(1.06), y, Inches(5.40), Inches(0.42), c, size=10,
         color=RGBColor(0x33, 0x3D, 0x49), anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.49)

box(s, Inches(6.82), Inches(1.40), Inches(5.90), Inches(4.90), fill=WHITE,
    line=RGBColor(0xD3, 0xDC, 0xE8))
text(s, Inches(7.06), Inches(1.52), Inches(5.4), Inches(0.30),
     "Deployment facts", size=14, color=NAVY, bold=True)
facts = [
    ("Base image", "ros:<distro>-ros-base"),
    ("Network", "host \u2014 DDS discovery is multicast"),
    ("Workspace", "src/ bind-mounted to /workspace/src"),
    ("GPU order", "nvidia runtime \u2192 /dev/dri \u2192 llvmpipe"),
    ("Display", "X11 preferred (RViz/Gazebo need GLX)"),
    ("Anti-flicker", "vblank_mode=xpresent"),
    ("Log rotation", "json-file, 50 MB \u00d7 5"),
    ("Ports", "8765 (Foxglove WebSocket)"),
]
y = Inches(1.92)
for n, v in facts:
    box(s, Inches(7.06), y, Inches(5.42), Inches(0.48), fill=LIGHT)
    text(s, Inches(7.20), y, Inches(2.0), Inches(0.48), n, size=10.5,
         color=BLUE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(9.26), y, Inches(3.14), Inches(0.48), v, size=10,
         color=RGBColor(0x33, 0x3D, 0x49), anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.55)
footer(s, "Kill order matters: watchdog and monitor die BEFORE docker rm -f, "
          "or the watchdog relaunches the container being removed.")
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "The kill order is a real bug that was fixed: removing the container "
    "while the watchdog was alive triggered an unwanted relaunch."
)

# ================================================================ 16 TAKEAWAY
s = blank()
box(s, Emu(0), Emu(0), SW, SH, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
text(s, Inches(1.05), Inches(0.66), Inches(11.2), Inches(0.66),
     "Engineering takeaways", size=32, color=WHITE, bold=True)

pts = [
    ("1", "Contracts over coupling",
     "Four interfaces (one action, one service, two topics) let the same "
     "supervisor drive a simulator or hardware.", GREEN),
    ("2", "Every wait needs a bound and a verdict",
     "Timeouts, detectors and a total-return budget turned unbounded stalls "
     "into 15\u201345 s skips.", ORANGE),
    ("3", "Measure, then tune",
     "Tolerances, SLAM travel gates and watchdog hysteresis were all set from "
     "observed failures \u2014 not defaults.", BLUE),
]
y = Inches(1.62)
for n, t, b, c in pts:
    box(s, Inches(1.05), y, Inches(11.2), Inches(1.30), fill=RGBColor(0x1D, 0x3A, 0x66),
        line=RGBColor(0x2E, 0x55, 0x8C))
    box(s, Inches(1.32), y + Inches(0.31), Inches(0.68), Inches(0.68),
        fill=c, shape=MSO_SHAPE.OVAL)
    text(s, Inches(1.32), y + Inches(0.31), Inches(0.68), Inches(0.68), n,
         size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(s, Inches(2.25), y + Inches(0.16), Inches(9.7), Inches(0.44), t,
         size=18, color=WHITE, bold=True)
    text(s, Inches(2.25), y + Inches(0.68), Inches(9.7), Inches(0.50), b,
         size=13, color=RGBColor(0xB6, 0xCB, 0xE6))
    y += Inches(1.46)

text(s, Inches(1.05), Inches(6.06), Inches(11.2), Inches(0.36),
     [("Full detail: ", {"size": 13}),
      ("DOCUMENTATION.md", {"size": 13, "font": MONO}),
      ("  \u00b7  ", {"size": 13}),
      ("CONTEXT.md", {"size": 13, "font": MONO}),
      (" (decision ledger)", {"size": 13})],
     color=RGBColor(0x8F, 0xB4, 0xDC))
text(s, Inches(1.05), Inches(6.50), Inches(11.2), Inches(0.36),
     [("Code: ", {"size": 13}),
      ("src/house_cleaner_mission/mission_supervisor.py",
       {"size": 13, "font": MONO}),
      ("  (1193 lines)", {"size": 13})],
     color=RGBColor(0x8F, 0xB4, 0xDC))
text(s, Inches(1.05), Inches(6.94), Inches(11.2), Inches(0.36),
     "Questions welcome \u2014 happy to walk the code or run the stack live.",
     size=13, color=RGBColor(0x8F, 0xB4, 0xDC), italic=True)
prs.slides[-1].notes_slide.notes_text_frame.text = (
    "Close on the three engineering principles, then point at "
    "DOCUMENTATION.md for depth."
)

# ---------------------------------------------------------------- save
here = os.path.dirname(os.path.abspath(__file__))
out_dir = os.path.join(os.path.dirname(here), "docs")
os.makedirs(out_dir, exist_ok=True)
out = os.path.join(out_dir, "House_Cleaner_Robot_Technical.pptx")
prs.save(out)
print(f"Wrote {out} ({len(prs.slides._sldIdLst)} slides)")
